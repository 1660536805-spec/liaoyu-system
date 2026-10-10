# -*- coding: utf-8 -*-
"""
把 outputs/aic-report/report.html 转成符合大赛模板的 .docx
- A4 纵向；页边距 上下 2.5cm / 左右 3cm
- 页眉：左上角官方 AIC 标志（取自模板 .doc 提取的 PNG）+ 居中赛事名
- 页脚：第 X 页 / 共 Y 页（PAGE / NUMPAGES 域）
- 标题层级：一级 三号(16pt) / 二级 四号(14pt) / 三级 小四(12pt)，均加粗
- 正文 小四(12pt)，单倍行距
- 目录：Word TOC 域（带缓存结果，打开后自动/一键更新）
所有手工插入的 OOXML 元素都按 ECMA-376 的子元素顺序插入，避免 Word 报“文件已损坏”。
用法: python gen_docx.py <output.docx>
"""
import os
import re
import sys

from lxml import html as lxml_html

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
# 用法: python gen_docx.py <out.docx> [in.html]
HTML = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, 'report.html')
LOGO = os.path.join(HERE, 'assets', 'aic-logo.png')
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'out.docx')

SERIF = '宋体'
SANS = 'Times New Roman'   # 西文与数字
MONO = 'Consolas'
ACCENT = '8A3B1E'          # 赭石
RACE_TITLE = '2026年第八届全球校园人工智能算法精英大赛'
TEXT_W = 15.0              # 正文栏宽 cm = 21 - 3 - 3

# ---------------------------------------------------------------- schema 顺序
PPR_SEQ = (
    'w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr',
    'w:widowControl', 'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd',
    'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
    'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
    'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
    'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
    'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl',
    'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange',
)
RPR_SEQ = (
    'w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps',
    'w:smallCaps', 'w:strike', 'w:dstrike', 'w:outline', 'w:shadow', 'w:emboss',
    'w:imprint', 'w:noProof', 'w:snapToGrid', 'w:vanish', 'w:webHidden',
    'w:color', 'w:spacing', 'w:w', 'w:kern', 'w:position', 'w:sz', 'w:szCs',
    'w:highlight', 'w:u', 'w:effect', 'w:bdr', 'w:shd', 'w:fitText',
    'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang', 'w:eastAsianLayout',
    'w:specVanish', 'w:oMath',
)
TBLPR_SEQ = (
    'w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual',
    'w:tblStyleRowBandSize', 'w:tblStyleColBandSize', 'w:tblW', 'w:jc',
    'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd', 'w:tblLayout',
    'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription',
    'w:tblPrChange',
)
TCPR_SEQ = (
    'w:cnfStyle', 'w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge', 'w:tcBorders',
    'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign',
    'w:hideMark', 'w:headers', 'w:cellIns', 'w:cellDel', 'w:cellMerge',
    'w:tcPrChange',
)
# settings.xml 中 updateFields 必须位于这些元素之前
SETTINGS_AFTER = ('w:compat', 'w:rsids', 'w:docVars', 'w:endnotePr',
                  'w:footnotePr', 'w:hdrShapeDefaults')


def insert_ordered(parent, el, seq):
    """把 el 插到 parent 中符合 seq 顺序的位置（不存在则 append）"""
    tags = [qn(t) for t in seq]
    try:
        idx = tags.index(el.tag)
    except ValueError:
        parent.append(el)
        return el
    for child in parent:
        try:
            ci = tags.index(child.tag)
        except ValueError:
            continue
        if ci > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


# ---------------------------------------------------------------- 基础工具
def fmt_run(run, size=12.0, bold=False, mono=False, color=None,
            underline=False, spacing=None):
    name_ascii = MONO if mono else SANS
    run.font.name = name_ascii                      # 走 python-docx，顺序天然正确
    rPr = run._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    rf.set(qn('w:ascii'), name_ascii)
    rf.set(qn('w:hAnsi'), name_ascii)
    rf.set(qn('w:eastAsia'), SERIF)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.underline = underline
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if spacing is not None:
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:val'), str(spacing))
        insert_ordered(rPr, sp, RPR_SEQ)
    return run


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        insert_ordered(tcPr, shd, TCPR_SEQ)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)


def set_tbl_borders(table, color='D9CFB8', sz=4, left=None):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        if edge == 'left' and left:
            e.set(qn('w:sz'), str(left[0]))
            e.set(qn('w:color'), left[1])
        else:
            e.set(qn('w:sz'), str(sz))
            e.set(qn('w:color'), color)
        e.set(qn('w:space'), '0')
        b.append(e)
    insert_ordered(tblPr, b, TBLPR_SEQ)


def fix_layout(table, widths):
    table.autofit = False                            # 生成 w:tblLayout
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = insert_ordered(tblPr, OxmlElement('w:tblLayout'), TBLPR_SEQ)
    layout.set(qn('w:type'), 'fixed')
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def add_field(par, instr, size=9.0):
    """插入一个 Word 域，如 PAGE / NUMPAGES"""
    r1 = par.add_run()
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fc)

    r2 = par.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = instr
    r2._r.append(it)

    r3 = par.add_run()
    fc2 = OxmlElement('w:fldChar')
    fc2.set(qn('w:fldCharType'), 'end')
    r3._r.append(fc2)
    for r in (r1, r2, r3):
        fmt_run(r, size=size)
    return par


def set_para_bottom_border(p, color='B9A98C', sz=4):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr')
        insert_ordered(pPr, bdr, PPR_SEQ)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), color)
    bdr.append(bottom)


# ---------------------------------------------------------------- 富文本渲染
def _add_text(par, txt, **kw):
    txt = txt.replace('\u00a0', ' ')
    txt = re.sub(r'\s*\n\s*', ' ', txt)
    if txt == '':
        return
    fmt_run(par.add_run(txt), **kw)


def render_inline(par, node, **kw):
    kw.setdefault('size', 12.0)
    kw.setdefault('bold', False)
    kw.setdefault('mono', False)
    kw.setdefault('color', None)

    if node.tag == 'br':
        par.add_run().add_break()
        return
    if node.text:
        _add_text(par, node.text, **kw)
    for ch in node:
        ckw = dict(kw)
        cls = ch.get('class') or ''
        if ch.tag == 'em':
            ckw['bold'] = True
            ckw['color'] = ACCENT
        elif ch.tag in ('strong', 'b'):
            ckw['bold'] = True
        elif ch.tag == 'code':
            ckw['mono'] = True
            ckw['size'] = kw['size'] - 1.0
        elif ch.tag == 'span' and 'kv' in cls:
            ckw['bold'] = True
        elif ch.tag == 'span' and 'mono' in cls:
            ckw['mono'] = True
        render_inline(par, ch, **ckw)
        if ch.tail:
            _add_text(par, ch.tail, **kw)


def plain_text(node):
    return re.sub(r'\s+', ' ', ''.join(node.itertext())).strip()


# ---------------------------------------------------------------- 文档骨架
def build_doc():
    doc = Document()

    st = doc.styles['Normal']
    st.font.name = SANS
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.0

    heads = {'Heading 1': (16, 14, 8), 'Heading 2': (14, 11, 6), 'Heading 3': (12, 9, 4)}
    for name, (size, before, after) in heads.items():
        s = doc.styles[name]
        s.font.name = SANS
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.line_spacing = 1.0
        s.paragraph_format.keep_with_next = True

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(3.0)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)

    # 打开时自动更新域（目录页码自动刷新）
    settings = doc.settings.element
    old = settings.find(qn('w:updateFields'))
    if old is not None:
        settings.remove(old)
    upd = OxmlElement('w:updateFields')
    upd.set(qn('w:val'), 'true')
    anchor = None
    for tag in SETTINGS_AFTER:
        anchor = settings.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(upd)
    else:
        settings.append(upd)

    build_header(sec)
    build_footer(sec)
    return doc


def build_header(sec):
    hdr = sec.header
    hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)

    p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.space_after = Pt(2)

    p.add_run().add_picture(LOGO, height=Cm(0.8))
    fmt_run(p.add_run('\t' + RACE_TITLE), size=9.0, color='333333')
    set_para_bottom_border(p)


def build_footer(sec):
    ftr = sec.footer
    ftr.is_linked_to_previous = False
    p = ftr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_run(p.add_run('第 '), size=9.0, color='333333')
    add_field(p, ' PAGE ')
    fmt_run(p.add_run(' 页 / 共 '), size=9.0, color='333333')
    add_field(p, ' NUMPAGES ')
    fmt_run(p.add_run(' 页'), size=9.0, color='333333')


# ---------------------------------------------------------------- 封面
def build_cover(doc, cover):
    def line(text, size, bold=False, color=None, before=0, spacing=None, after=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.0
        fmt_run(p.add_run(text), size=size, bold=bold, color=color, spacing=spacing)
        return p

    line('2026年第八届', 15, before=30)
    line('全球校园人工智能', 26, bold=True, before=18, spacing=60)
    line('算法精英大赛', 26, bold=True, spacing=60, after=10)
    line('算法创新赛', 17, bold=True, before=6)
    line(cover.get('data-track') or '（AI+软件赛道）', 15, bold=True, color=ACCENT, before=4)
    line(cover.get('data-type') or '技 术 报 告', 22, bold=True, before=44, spacing=200)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(46)
    sp.paragraph_format.space_after = Pt(0)

    def meta(label, value, pad_to=12):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(10)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(2.6)
        fmt_run(p.add_run(label), size=14)
        w = sum(1.0 if ord(c) > 0x2E80 else 0.5 for c in value)
        pad = '\u3000' * max(1, int(round(pad_to - w)))
        fmt_run(p.add_run(value + pad), size=14, underline=True)

    meta('团队名称：', '锻五音')
    meta('参赛编号：', 'AIC-2026-')
    meta('作品名称：', '弦养')

    d = line(cover.get('data-date') or '2026年10月10日', 14, before=170)
    d.add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- 目录
def build_toc(doc, toc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(16)
    fmt_run(p.add_run('目    录'), size=22, bold=True, spacing=160)

    lines = toc.xpath(".//div[contains(@class,'toc-line')]")
    instr = ' TOC \\o "1-3" \\h \\z \\u '

    for i, ln in enumerate(lines):
        cls = ln.get('class') or ''
        level = 1 if 'toc-l1' in cls else (2 if 'toc-l2' in cls else 3)
        t = plain_text(ln.xpath(".//span[@class='t']")[0])
        pg_node = ln.xpath(".//span[@class='pg']")
        pg = plain_text(pg_node[0]) if pg_node else ''

        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.space_before = Pt(2 if level == 1 else 0)
        pf.space_after = Pt(1)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0 if level == 1 else (0.75 if level == 2 else 1.5))
        pf.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

        if i == 0:
            r1 = par.add_run()
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), 'begin')
            fc.set(qn('w:dirty'), 'true')
            r1._r.append(fc)
            r2 = par.add_run()
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = instr
            r2._r.append(it)
            r3 = par.add_run()
            fc2 = OxmlElement('w:fldChar')
            fc2.set(qn('w:fldCharType'), 'separate')
            r3._r.append(fc2)

        fmt_run(par.add_run(t), size=11, bold=(level == 1))
        fmt_run(par.add_run('\t'), size=11)
        fmt_run(par.add_run(pg), size=11, bold=(level == 1))

        if i == len(lines) - 1:
            r4 = par.add_run()
            fc3 = OxmlElement('w:fldChar')
            fc3.set(qn('w:fldCharType'), 'end')
            r4._r.append(fc3)

    b = doc.add_paragraph()
    b.add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- 正文块
def add_heading(doc, node):
    # HTML 的 h2/h3/h4 → 模板的一级(三号)/二级(四号)/三级(小四)标题
    lvl = min(int(node.tag[1]) - 1, 3)
    p = doc.add_paragraph(style='Heading %d' % lvl)
    if lvl == 1:
        # 与 PDF 版式一致：带 class="sec" 的一级标题（各章）另起一页
        cls = (node.get('class') or '').split()
        p.paragraph_format.page_break_before = ('sec' in cls) or (plain_text(node) == '一、项目概述')
    render_inline(p, node, size={1: 16.0, 2: 14.0, 3: 12.0}[lvl], bold=True)
    for r in p.runs:
        r.font.bold = True
    return p


def add_para(doc, node):
    cls = node.get('class') or ''
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if 'cap' in cls:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(4)
        pf.space_after = Pt(4)
        render_inline(p, node, size=10.5, bold=True)
    elif 'sub' in cls:
        pf.space_before = Pt(8)
        pf.space_after = Pt(2)
        render_inline(p, node, size=12, bold=True)
        for r in p.runs:
            r.font.bold = True
    elif 'note' in cls:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        render_inline(p, node, size=10.5, color='444444')
    elif 'mono' in cls:
        pf.space_before = Pt(3)
        pf.space_after = Pt(5)
        render_inline(p, node, size=10.5, mono=True)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        render_inline(p, node, size=12)
    return p


def add_list(doc, ul):
    size = 12.0
    for li in ul.iterchildren('li'):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.8)
        pf.first_line_indent = Cm(-0.4)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt_run(p.add_run('· '), size=size)
        render_inline(p, li, size=size)


def add_pre(doc, node):
    txt = ''.join(node.itertext())
    lines = txt.split('\n')
    while lines and not lines[-1].strip():
        lines.pop()
    while lines and not lines[0].strip():
        lines.pop(0)

    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_tbl_borders(t, color='D9CFB8', sz=4, left=(18, 'B0552E'))
    fix_layout(t, [TEXT_W])
    cell = t.cell(0, 0)
    shade(cell, 'F7F3EA')

    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.0
    for i, ln in enumerate(lines):
        if i:
            p.add_run().add_break()
        fmt_run(p.add_run(ln), size=9.0, mono=True)


def add_box(doc, node):
    t = doc.add_table(rows=1, cols=1)
    set_tbl_borders(t, color='C9B9A0', sz=4)
    fix_layout(t, [TEXT_W])
    cell = t.cell(0, 0)
    shade(cell, 'FBF7EF')

    first = True
    for k in [c for c in node.iterchildren() if c.tag == 'p']:
        cls = k.get('class') or ''
        if first:
            p = cell.paragraphs[0]
            first = False
        else:
            p = cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        if 'note' in cls:
            render_inline(p, k, size=10.5, color='444444')
        elif 'mono' in cls:
            render_inline(p, k, size=10.5, mono=True)
        else:
            render_inline(p, k, size=12)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_widths(cells):
    pcts = []
    for c in cells:
        st = c.get('style') or ''
        m = re.search(r'width\s*:\s*([\d.]+)%', st)
        pcts.append(float(m.group(1)) if m else None)
    unknown = [i for i, v in enumerate(pcts) if v is None]
    if unknown:
        rest = max(0.0, 100.0 - sum(v for v in pcts if v)) / len(unknown)
        vals = [v if v else rest for v in pcts]
    else:
        s = sum(pcts) or 100.0
        vals = [v / s * 100.0 for v in pcts]
    return [v / 100.0 * TEXT_W for v in vals]


def add_table(doc, node):
    rows = node.findall('.//tr')
    if not rows:
        return
    cells_per_row = [[c for c in tr if c.tag in ('td', 'th')] for tr in rows]
    ncols = max(len(r) for r in cells_per_row)
    widths = parse_widths(cells_per_row[0]) if cells_per_row[0] else None
    if not widths or len(widths) != ncols:
        widths = [TEXT_W / ncols] * ncols

    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = 'Table Grid'
    fix_layout(t, widths)

    for ri, cells in enumerate(cells_per_row):
        for ci in range(ncols):
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            if ci >= len(cells):
                continue
            src = cells[ci]
            cls = src.get('class') or ''
            header = src.tag == 'th'
            if header:
                shade(cell, 'F0E9DC')
            if 'c' in cls.split() or header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            render_inline(p, src, size=10.5, bold=header)
            if header:
                for r in p.runs:
                    r.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ---------------------------------------------------------------- 主流程
def main():
    with open(HTML, 'rb') as f:
        root = lxml_html.fromstring(f.read())
    body = root.find('body')

    doc = build_doc()

    for node in body.iterchildren():
        tag = node.tag
        cls = node.get('class') or ''
        if tag == 'div' and 'cover' in cls:
            build_cover(doc, node)
        elif tag == 'div' and 'toc' in cls:
            build_toc(doc, node)
        elif tag == 'div' and 'box' in cls:
            add_box(doc, node)
        elif tag in ('h2', 'h3', 'h4'):
            add_heading(doc, node)
        elif tag == 'p':
            if not plain_text(node):
                continue
            add_para(doc, node)
        elif tag == 'ul':
            add_list(doc, node)
        elif tag == 'pre':
            add_pre(doc, node)
        elif tag == 'table':
            add_table(doc, node)

    doc.save(OUT)
    print('DOCX ->', OUT, os.path.getsize(OUT), 'bytes')


if __name__ == '__main__':
    main()
